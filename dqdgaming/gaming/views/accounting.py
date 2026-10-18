"""Admin accounting dashboards and exports (Excel, PDF, Google Sheets)."""

from datetime import date, datetime, time, timedelta
from decimal import Decimal

import gspread
import openpyxl
from django.conf import settings
from django.db.models import Avg, Count, Sum
from django.db.models.functions import ExtractYear, TruncDate, TruncMonth, TruncWeek
from django.http import HttpResponse
from django.utils import timezone
from google.oauth2.service_account import Credentials
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from gaming.models import Booking
from gaming.permissions import IsAdminUserRole


def get_booking_queryset(request):

    queryset = Booking.objects.filter(is_deleted=False)

    period = request.GET.get(
        "period",
        "monthly",
    )

    today = timezone.localdate()

    if period == "daily":

        queryset = queryset.filter(booking_date=today)

    elif period == "weekly":

        week_start = today - timedelta(days=today.weekday())

        queryset = queryset.filter(booking_date__gte=week_start)

    elif period == "monthly":

        queryset = queryset.filter(
            booking_date__month=today.month,
            booking_date__year=today.year,
        )

    elif period == "yearly":

        queryset = queryset.filter(
            booking_date__year=today.year,
        )

    elif period == "custom":

        start = request.GET.get("start_date")

        end = request.GET.get("end_date")

        if start and end:

            queryset = queryset.filter(
                booking_date__range=[
                    start,
                    end,
                ]
            )

    return queryset


class AdminAccountingDashboardAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        bookings = get_booking_queryset(request)

        completed = bookings.filter(status=Booking.Status.COMPLETED)

        total_bookings = bookings.count()

        completed_bookings = completed.count()

        cancelled = bookings.filter(status=Booking.Status.CANCELLED).count()

        pending = bookings.filter(status=Booking.Status.PENDING).count()

        gross = completed.aggregate(total=Sum("subtotal"))["total"] or 0

        discount = completed.aggregate(total=Sum("discount_amount"))["total"] or 0

        revenue = completed.aggregate(total=Sum("total_amount"))["total"] or 0

        received = completed.aggregate(total=Sum("price_paid"))["total"] or 0

        average = completed.aggregate(avg=Avg("total_amount"))["avg"] or 0

        pending_payment = (
            bookings.filter(payment_status=Booking.PaymentStatus.PENDING).aggregate(
                total=Sum("total_amount")
            )["total"]
            or 0
        )

        return Response(
            {
                "summary": {
                    "total_bookings": total_bookings,
                    "completed_bookings": completed_bookings,
                    "cancelled_bookings": cancelled,
                    "pending_bookings": pending,
                    "gross_revenue": gross,
                    "discount": discount,
                    "net_revenue": revenue,
                    "amount_received": received,
                    "pending_payment": pending_payment,
                    "average_booking": average,
                    "profit": received,
                    "loss": discount,
                }
            }
        )


class AdminRevenueChartAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        period = request.GET.get("period", "monthly")

        queryset = get_booking_queryset(request).filter(status=Booking.Status.COMPLETED)

        if period == "daily":
            queryset = queryset.annotate(label=TruncDate("booking_date"))

        elif period == "weekly":
            queryset = queryset.annotate(label=TruncWeek("booking_date"))

        elif period == "monthly":
            queryset = queryset.annotate(label=TruncMonth("booking_date"))

        elif period == "yearly":
            queryset = queryset.annotate(label=ExtractYear("booking_date"))

        else:  # custom
            queryset = queryset.annotate(label=TruncDate("booking_date"))

        data = (
            queryset.values("label")
            .annotate(
                bookings=Count("id"),
                revenue=Sum("price_paid"),
                discount=Sum("discount_amount"),
            )
            .order_by("label")
        )

        results = []

        for row in data:
            results.append(
                {
                    "label": str(row["label"]),
                    "bookings": row["bookings"] or 0,
                    "revenue": float(row["revenue"] or 0),
                    "discount": float(row["discount"] or 0),
                }
            )

        return Response(results)


class AdminAccountingBookingsAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        bookings = (
            get_booking_queryset(request)
            .select_related(
                "user",
                "item",
            )
            .order_by("-booking_date")
        )

        data = []

        for booking in bookings:

            data.append(
                {
                    "booking_id": booking.booking_id,
                    "customer": (
                        booking.customer_name or booking.user.full_name
                        if booking.user
                        else "-"
                    ),
                    "booking_date": booking.created_at.date(),
                    "game": (booking.item.name if booking.item else "Deleted Game"),
                    "hours": booking.total_hours or 0,
                    "subtotal": booking.subtotal or 0,
                    "discount": booking.discount_amount or 0,
                    "total": booking.total_amount or 0,
                    "paid": booking.price_paid or 0,
                    "payment_status": booking.payment_status,
                    "booking_status": booking.status,
                }
            )

        return Response(data)


class AdminProfitLossAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        bookings = get_booking_queryset(request)

        completed = bookings.filter(status=Booking.Status.COMPLETED)

        gross_income = completed.aggregate(total=Sum("subtotal"))["total"] or 0

        discount = completed.aggregate(total=Sum("discount_amount"))["total"] or 0

        revenue = completed.aggregate(total=Sum("price_paid"))["total"] or 0

        # Replace later with Expense model
        expenses = 0

        profit = revenue - expenses

        return Response(
            {
                "gross_income": gross_income,
                "discount": discount,
                "expenses": expenses,
                "net_revenue": revenue,
                "profit": profit,
                "loss": max(0, expenses - revenue),
            }
        )


class AdminAccountingExcelExportAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        workbook = openpyxl.Workbook()

        sheet = workbook.active

        sheet.title = "Accounting"

        headers = [
            "Booking ID",
            "Customer",
            "Game",
            "Date",
            "Start Time",
            "End Time",
            "Hours",
            "Subtotal",
            "Discount",
            "Total",
            "Payment Status",
            "Paid",
            "Notes",
            "Status",
        ]

        for i, h in enumerate(headers, start=1):
            sheet.cell(row=1, column=i).value = h

        bookings = get_booking_queryset(request)

        row = 2

        for booking in bookings:

            sheet.cell(row=row, column=1).value = booking.booking_id
            sheet.cell(row=row, column=2).value = booking.customer_name
            sheet.cell(row=row, column=3).value = (
                booking.item.name if booking.item else "-"
            )
            sheet.cell(row=row, column=4).value = str(booking.booking_date)
            sheet.cell(row=row, column=5).value = str(booking.start_time)
            sheet.cell(row=row, column=6).value = str(booking.end_time)

            sheet.cell(row=row, column=7).value = float(booking.total_hours)
            sheet.cell(row=row, column=8).value = float(booking.subtotal)
            sheet.cell(row=row, column=9).value = float(booking.discount_amount)
            sheet.cell(row=row, column=10).value = float(booking.total_amount)

            sheet.cell(row=row, column=11).value = str(booking.payment_status)
            sheet.cell(row=row, column=12).value = float(booking.price_paid)
            sheet.cell(row=row, column=13).value = str(booking.notes)

            sheet.cell(row=row, column=14).value = booking.status

            row += 1

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        response["Content-Disposition"] = 'attachment; filename="Accounting.xlsx"'

        workbook.save(response)

        return response


def sheet_value(value):
    if value is None:
        return ""

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, time):
        return value.strftime("%H:%M:%S")

    return value


class AdminAccountingPDFExportAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        response = HttpResponse(content_type="application/pdf")

        response["Content-Disposition"] = 'attachment; filename="Accounting.pdf"'

        doc = SimpleDocTemplate(
            response,
            pagesize=landscape(A4),
            leftMargin=10,
            rightMargin=10,
            topMargin=15,
            bottomMargin=15,
        )

        data = [
            [
                "Booking ID",
                "Customer",
                "Game / Item",
                "Booking Date",
                "Start Time",
                "End Time",
                "Total Hours",
                "Subtotal",
                "Discount",
                "Total",
                "Payment",
                "Paid",
                "Notes",
                "Status",
            ]
        ]

        bookings = get_booking_queryset(request)

        for booking in bookings:

            data.append(
                [
                    sheet_value(booking.booking_id),
                    sheet_value(booking.customer_name),
                    sheet_value(
                        booking.item.name
                        if booking.item
                        else booking.combo_pack.name if booking.combo_pack else ""
                    ),
                    sheet_value(booking.booking_date),
                    sheet_value(booking.start_time),
                    sheet_value(booking.end_time),
                    sheet_value(booking.total_hours),
                    sheet_value(booking.subtotal),
                    sheet_value(booking.discount_amount),
                    sheet_value(booking.total_amount),
                    sheet_value(booking.payment_status),
                    sheet_value(booking.price_paid),
                    sheet_value(booking.notes),
                    sheet_value(booking.status),
                ]
            )

        table = Table(
            data,
            repeatRows=1,
            colWidths=[
                55,  # Booking ID
                75,  # Customer
                70,  # Game
                55,  # Date
                40,  # Start
                40,  # End
                50,  # Hours
                50,  # Subtotal
                50,  # Discount
                50,  # Total
                50,  # Payment
                50,  # Paid
                120,  # Notes
                45,  # Status
            ],
        )

        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ]
            )
        )

        styles = getSampleStyleSheet()

        elements = [
            Paragraph("<b>DQD Gaming Accounting Report</b>", styles["Title"]),
            Spacer(1, 12),
            table,
        ]

        doc.build(elements)

        return response


class AdminAccountingGoogleSheetSyncAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request):

        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]

        credentials = Credentials.from_service_account_file(
            settings.GOOGLE_SERVICE_ACCOUNT_INFO,
            scopes=[
                "https://www.googleapis.com/auth/spreadsheets",
            ],
        )

        client = gspread.authorize(credentials)

        sheet = client.open_by_key(settings.GOOGLE_SHEET_ID).sheet1

        sheet.clear()

        rows = [
            [
                "Booking ID",
                "Customer",
                "Game / Item",
                "Booking Date",
                "Start Time",
                "End Time",
                "Total Hours",
                "Subtotal",
                "Discount",
                "Total",
                "Payment",
                "Paid",
                "Notes",
                "Status",
            ]
        ]

        bookings = Booking.objects.filter(is_deleted=False)

        for booking in bookings:

            rows.append(
                [
                    sheet_value(booking.booking_id),
                    sheet_value(booking.customer_name),
                    sheet_value(
                        booking.item.name
                        if booking.item
                        else booking.combo_pack.name if booking.combo_pack else ""
                    ),
                    sheet_value(booking.booking_date),
                    sheet_value(booking.start_time),
                    sheet_value(booking.end_time),
                    sheet_value(booking.total_hours),
                    sheet_value(booking.subtotal),
                    sheet_value(booking.discount_amount),
                    sheet_value(booking.total_amount),
                    sheet_value(booking.payment_status),
                    sheet_value(booking.price_paid),
                    sheet_value(booking.notes),
                    sheet_value(booking.status),
                ]
            )

        sheet.clear()

        sheet.update(
            values=rows,
            range_name="A1",
        )

        return Response(
            {
                "success": True,
                "message": "Google Sheet updated successfully.",
                "rows": len(rows) - 1,
            }
        )
