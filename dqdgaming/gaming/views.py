import logging
import os
from urllib import request
from urllib.parse import urlparse
from django.shortcuts import get_object_or_404
from rest_framework_simplejwt.settings import api_settings
from gaming.utils.emailjs import send_game_booking_email
from rest_framework import status
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView
import requests as http_requests
from django.conf import settings
from django.contrib.auth.password_validation import validate_password
from django.contrib.auth.tokens import default_token_generator
from django.core.exceptions import ValidationError as DjangoValidationError
from django.core.mail import send_mail
from django.db.models import Q, Sum
from django.utils import timezone
from django.utils.encoding import force_bytes, force_str
from django.utils.http import (
    urlsafe_base64_decode,
    urlsafe_base64_encode,
)
from google.auth.transport import requests as google_requests
from google.oauth2 import id_token as google_id_token
from rest_framework import generics, permissions, status
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.exceptions import TokenError
from rest_framework_simplejwt.tokens import RefreshToken

from gaming.tracking import (
    mark_current_device_offline,
    register_device,
)
from gaming.middleware import DeviceActivityMixin
from gaming.permissions import IsAdminUserRole

from gaming.models import (
    Booking,
    ComboPack,
    CustomUser,
    EventBooking,
    ExclusiveEvent,
    GamingItem,
    LoyaltyTransaction,
    SpinnerSpin,
    UserDevice,
)
from gaming.serializers import *

from gaming.utils.image_encryption import (
    encrypt_image_bytes,
    encrypt_uploaded_image,
    safe_image_to_data_uri,
)

from rest_framework.throttling import AnonRateThrottle, ScopedRateThrottle




logger = logging.getLogger(__name__)


# ============================================================
# AUTHENTICATION HELPERS
# ============================================================


def _build_token_response(
    user,
    request=None,
    is_new=False,
    include_is_new=False,
):
    data = {
        **get_tokens_for_user(user),
        "user": UserProfileSerializer(
            user,
            context={"request": request},
        ).data,
    }

    if include_is_new:
        data["is_new"] = is_new

    return data


def _role_redirect_path(user):
    if user.role == CustomUser.Role.ADMIN:
        return "/admin/dashboard"

    return "/user/dashboard"


def _persist_remote_profile_image(user, image_url):
    if not image_url or user.profile_image:
        return []

    try:
        response = http_requests.get(
            image_url,
            timeout=10,
            stream=True,
        )
        response.raise_for_status()
    except http_requests.RequestException as exc:
        logger.warning(
            "Could not fetch profile image for %s: %s",
            user.email,
            exc,
        )
        return []

    content_type = response.headers.get("Content-Type", "").lower()

    if not content_type.startswith("image/"):
        return []

    content = response.content

    if len(content) > 2 * 1024 * 1024:
        logger.warning(
            "Profile image for %s exceeded the 2MB limit.",
            user.email,
        )
        return []

    try:
        payload = encrypt_image_bytes(
            content=content,
            filename=os.path.basename(image_url.split("?", 1)[0]),
        )
    except DjangoValidationError as exc:
        logger.warning(
            "Profile image for %s could not be encrypted: %s",
            user.email,
            exc,
        )
        return []

    user.profile_image = payload.encrypted_content
    return ["profile_image"]


def _store_encrypted_qr_code(instance, qr_file):
    payload = encrypt_uploaded_image(qr_file)
    instance.qr_code = payload.encrypted_content
    return payload


def _finalize_login_response(
    user,
    request,
    is_new=False,
    include_is_new=False,
):
    validate_login_allowed(user)
    touch_last_login(user)

    response_data = _build_token_response(
        user,
        request=request,
        is_new=is_new,
        include_is_new=include_is_new,
    )
    response_data["redirect_to"] = _role_redirect_path(user)

    return response_data


def _finalize_device_login(
    user,
    request,
    *,
    is_new=False,
    include_is_new=False,
):
    response_data = _finalize_login_response(
        user,
        request,
        is_new=is_new,
        include_is_new=include_is_new,
    )

    device = register_device(request, user)
    response_data["device_id"] = str(device.device_id)

    return response_data


def _send_verification_email(user):
    uid = urlsafe_base64_encode(force_bytes(user.pk))
    token = default_token_generator.make_token(user)
    frontend_url = settings.FRONTEND_URL.rstrip("/")
    link = f"{frontend_url}/verify-email/{uid}/{token}/"

    send_mail(
        subject="Verify your DQD Gaming account",
        message=(
            f"Hi {user.first_name},\n\n"
            f"Verify your email using the following link:\n{link}"
        ),
        from_email=settings.DEFAULT_FROM_EMAIL,
        recipient_list=[user.email],
        fail_silently=False,
    )


# ============================================================
# REGISTRATION
# ============================================================


class RegisterView(APIView):
    permission_classes = [permissions.AllowAny]
    throttle_classes = [AnonRateThrottle, ScopedRateThrottle]
    throttle_scope = "register"

    def post(self, request):
        serializer = RegisterSerializer(
            data=request.data,
            context={"request": request},
        )
        serializer.is_valid(raise_exception=True)
        user = serializer.save()

        try:
            _send_verification_email(user)
        except Exception:
            logger.exception(
                "Verification email failed for %s.",
                user.email,
            )

        return Response(
            _finalize_device_login(user, request),
            status=status.HTTP_201_CREATED,
        )


# ============================================================
# EMAIL LOGIN
# ============================================================


class LoginView(APIView):
    permission_classes = [permissions.AllowAny]
    throttle_classes = [AnonRateThrottle, ScopedRateThrottle]
    throttle_scope = "login"

    def post(self, request):
        serializer = LoginSerializer(
            data=request.data,
            context={"request": request},
        )
        serializer.is_valid(raise_exception=True)

        user = serializer.validated_data["user"]
        response_data = _finalize_device_login(user, request)

        if not user.is_verified:
            response_data["email_unverified"] = True

        return Response(
            response_data,
            status=status.HTTP_200_OK,
        )


# ============================================================
# LOGOUT
# ============================================================


class LogoutView(DeviceActivityMixin, APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        refresh_token = request.data.get("refresh")
        if not refresh_token:
            return Response(
                {"detail": "Refresh token is required."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            refresh = RefreshToken(refresh_token)
            token_user_id = refresh.payload.get(api_settings.USER_ID_CLAIM)
            if str(token_user_id) != str(request.user.pk):
                raise TokenError("Refresh token does not belong to this user.")

            refresh.blacklist()
        except TokenError:
            return Response(
                {"detail": "Invalid or expired refresh token."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        mark_current_device_offline(request)
        return Response(
            {"detail": "Successfully logged out."},
            status=status.HTTP_200_OK,
        )


# ============================================================
# USER DEVICES
# ============================================================


class MyDevicesView(DeviceActivityMixin, APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        devices = UserDevice.objects.filter(
            user=request.user,
            is_deleted=False,
        ).order_by("-last_seen")

        serializer = UserDeviceSerializer(
            devices,
            many=True,
            context={"request": request},
        )

        return Response(
            serializer.data,
            status=status.HTTP_200_OK,
        )


# ============================================================
# PROFILE
# ============================================================


class ProfileView(
    DeviceActivityMixin,
    generics.RetrieveUpdateAPIView,
):
    serializer_class = UserProfileSerializer
    permission_classes = [permissions.IsAuthenticated]
    http_method_names = ["get", "patch"]

    def get_object(self):
        return self.request.user


# ============================================================
# PASSWORD CHANGE
# ============================================================


class PasswordChangeView(DeviceActivityMixin, APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        serializer = PasswordChangeSerializer(
            data=request.data,
            context={"request": request},
        )
        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(
            {"detail": "Password updated successfully."},
            status=status.HTTP_200_OK,
        )


# ============================================================
# PASSWORD RESET REQUEST
# ============================================================


class PasswordResetRequestView(APIView):
    permission_classes = [permissions.AllowAny]
    throttle_classes = [AnonRateThrottle, ScopedRateThrottle]
    throttle_scope = "password_reset"

    def post(self, request):
        email = request.data.get("email", "").strip().lower()

        if email:
            user = CustomUser.objects.filter(
                email__iexact=email,
                is_active=True,
                is_deleted=False,
            ).first()

            if user:
                uid = urlsafe_base64_encode(force_bytes(user.pk))
                token = default_token_generator.make_token(user)
                frontend_url = settings.FRONTEND_URL.rstrip("/")
                link = f"{frontend_url}/reset-password/" f"{uid}/{token}/"

                try:
                    send_mail(
                        subject="Reset your DQD Gaming password",
                        message=(
                            "Reset your password using the following " f"link:\n{link}"
                        ),
                        from_email=settings.DEFAULT_FROM_EMAIL,
                        recipient_list=[user.email],
                        fail_silently=False,
                    )
                except Exception:
                    logger.exception(
                        "Password reset email failed for %s.",
                        user.email,
                    )

        return Response(
            {
                "detail": (
                    "If that email is registered, you will receive "
                    "a reset link shortly."
                )
            },
            status=status.HTTP_200_OK,
        )


# ============================================================
# PASSWORD RESET CONFIRM
# ============================================================


class PasswordResetConfirmView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        uid = request.data.get("uid", "")
        token = request.data.get("token", "")
        new_password = request.data.get("new_password", "")

        try:
            user_id = force_str(urlsafe_base64_decode(uid))
            user = CustomUser.objects.get(
                pk=user_id,
                is_active=True,
                is_deleted=False,
            )
        except (
            CustomUser.DoesNotExist,
            ValueError,
            TypeError,
            OverflowError,
        ):
            return Response(
                {"detail": "Invalid reset link."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not default_token_generator.check_token(user, token):
            return Response(
                {"detail": "Reset link is invalid or has expired."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            validate_password(new_password, user=user)
        except DjangoValidationError as exc:
            return Response(
                {"new_password": list(exc.messages)},
                status=status.HTTP_400_BAD_REQUEST,
            )

        user.set_password(new_password)
        user.save(update_fields=["password", "updated_at"])

        UserDevice.objects.filter(
            user=user,
            is_online=True,
        ).update(
            is_online=False,
            logged_out_at=timezone.now(),
        )

        return Response(
            {"detail": "Password has been reset successfully."},
            status=status.HTTP_200_OK,
        )


# ============================================================
# EMAIL VERIFICATION
# ============================================================


class EmailVerifyView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        uid = request.data.get("uid", "")
        token = request.data.get("token", "")

        try:
            user_id = force_str(urlsafe_base64_decode(uid))
            user = CustomUser.objects.get(
                pk=user_id,
                is_deleted=False,
            )
        except (
            CustomUser.DoesNotExist,
            ValueError,
            TypeError,
            OverflowError,
        ):
            return Response(
                {"detail": "Invalid verification link."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not default_token_generator.check_token(user, token):
            return Response(
                {"detail": ("Verification link is invalid or has expired.")},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if not user.is_verified:
            user.is_verified = True
            user.save(update_fields=["is_verified", "updated_at"])

        return Response(
            {"detail": "Email verified successfully."},
            status=status.HTTP_200_OK,
        )


class EmailResendView(APIView):
    permission_classes = [permissions.AllowAny]

    def post(self, request):
        email = request.data.get("email", "").strip().lower()

        user = CustomUser.objects.filter(
            email__iexact=email,
            is_active=True,
            is_deleted=False,
            is_verified=False,
        ).first()

        if user:
            try:
                _send_verification_email(user)
            except Exception:
                logger.exception(
                    "Verification email resend failed for %s.",
                    user.email,
                )

        return Response(
            {
                "detail": (
                    "If that email is pending verification, "
                    "a new link has been sent."
                )
            },
            status=status.HTTP_200_OK,
        )


# ============================================================
# GOOGLE LOGIN
# ============================================================


class GoogleLoginView(APIView):
    permission_classes = [permissions.AllowAny]
    throttle_classes = [AnonRateThrottle, ScopedRateThrottle]
    throttle_scope = "social_login"

    def post(self, request):
        serializer = GoogleLoginSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        return _complete_google_login(
            request,
            serializer.validated_data["id_token"],
        )


class GoogleCodeLoginView(APIView):
    permission_classes = [permissions.AllowAny]
    throttle_classes = [AnonRateThrottle, ScopedRateThrottle]
    throttle_scope = "social_login"

    def post(self, request):
        serializer = GoogleCodeLoginSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        if not settings.GOOGLE_CLIENT_SECRET:
            return Response(
                {"detail": "Google redirect sign-in is not configured."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        redirect_uri = serializer.validated_data["redirect_uri"]
        parsed_redirect = urlparse(redirect_uri)
        redirect_origin = f"{parsed_redirect.scheme}://{parsed_redirect.netloc}"
        allowed_origins = {
            origin.rstrip("/") for origin in settings.CORS_ALLOWED_ORIGINS
        }

        if (
            redirect_origin not in allowed_origins
            or parsed_redirect.path != "/sign-in"
            or parsed_redirect.query
            or parsed_redirect.fragment
            or parsed_redirect.username
            or parsed_redirect.password
        ):
            return Response(
                {"detail": "Invalid Google redirect URI."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            token_response = http_requests.post(
                "https://oauth2.googleapis.com/token",
                data={
                    "code": serializer.validated_data["code"],
                    "client_id": settings.GOOGLE_CLIENT_ID,
                    "client_secret": settings.GOOGLE_CLIENT_SECRET,
                    "redirect_uri": redirect_uri,
                    "grant_type": "authorization_code",
                },
                timeout=10,
            )
            token_response.raise_for_status()
            token_data = token_response.json()
        except (http_requests.RequestException, ValueError):
            logger.warning("Google authorization-code exchange failed.")
            return Response(
                {"detail": "Google sign-in could not be verified."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        id_token = token_data.get("id_token")
        if not id_token:
            return Response(
                {"detail": "Google sign-in did not return an ID token."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        return _complete_google_login(request, id_token)


def _complete_google_login(request, token):
    try:
        google_data = google_id_token.verify_oauth2_token(
            token,
            google_requests.Request(),
            settings.GOOGLE_CLIENT_ID,
        )
    except ValueError as exc:
        logger.warning(
            "Google token verification failed: %s",
            exc,
        )
        return Response(
            {"detail": "Invalid Google token."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    email = google_data.get("email", "").strip().lower()
    first_name = google_data.get("given_name", "").strip()
    last_name = google_data.get("family_name", "").strip()
    picture = google_data.get("picture", "")
    google_email_verified = google_data.get(
        "email_verified",
        False,
    )

    if not email or not google_email_verified:
        return Response(
            {"detail": "Google account email is not verified."},
            status=status.HTTP_400_BAD_REQUEST,
        )

    user, is_new = CustomUser.objects.get_or_create(
        email=email,
        defaults={
            "first_name": first_name or "Google User",
            "last_name": last_name,
            "login_provider": CustomUser.LoginProvider.GOOGLE,
            "is_verified": True,
            "role": CustomUser.Role.USER,
        },
    )

    validate_login_allowed(user)

    update_fields = []

    if not user.first_name and first_name:
        user.first_name = first_name
        update_fields.append("first_name")

    if not user.last_name and last_name:
        user.last_name = last_name
        update_fields.append("last_name")

    if not user.is_verified:
        user.is_verified = True
        update_fields.append("is_verified")

    update_fields.extend(_persist_remote_profile_image(user, picture))

    if update_fields:
        update_fields.append("updated_at")
        user.save(update_fields=list(dict.fromkeys(update_fields)))

    return Response(
        _finalize_device_login(
            user,
            request,
            is_new=is_new,
            include_is_new=True,
        ),
        status=status.HTTP_200_OK,
    )


# ============================================================
# FACEBOOK LOGIN
# ============================================================


class FacebookLoginView(APIView):
    permission_classes = [permissions.AllowAny]
    throttle_classes = [AnonRateThrottle, ScopedRateThrottle]
    throttle_scope = "social_login"

    FB_GRAPH_URL = "https://graph.facebook.com/me"
    FB_FIELDS = "id,name,first_name,last_name,email,picture"

    def post(self, request):
        serializer = FacebookLoginSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        access_token = serializer.validated_data["access_token"]

        try:
            response = http_requests.get(
                self.FB_GRAPH_URL,
                params={
                    "fields": self.FB_FIELDS,
                    "access_token": access_token,
                },
                timeout=10,
            )
            response.raise_for_status()
            facebook_data = response.json()
        except (
            http_requests.RequestException,
            ValueError,
        ) as exc:
            logger.warning(
                "Facebook Graph API request failed: %s",
                exc,
            )
            return Response(
                {"detail": "Could not validate Facebook token."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        if "error" in facebook_data:
            return Response(
                {
                    "detail": facebook_data["error"].get(
                        "message",
                        "Facebook authentication failed.",
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        email = facebook_data.get("email", "").strip().lower()
        first_name = facebook_data.get(
            "first_name",
            facebook_data.get("name", ""),
        ).strip()
        last_name = facebook_data.get("last_name", "").strip()

        picture = facebook_data.get("picture", {}).get("data", {}).get("url", "")

        if not email:
            return Response(
                {
                    "detail": (
                        "Facebook did not provide an email address. "
                        "Ensure the account has a verified email."
                    )
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        user, is_new = CustomUser.objects.get_or_create(
            email=email,
            defaults={
                "first_name": first_name or "Facebook User",
                "last_name": last_name,
                "login_provider": CustomUser.LoginProvider.FACEBOOK,
                "is_verified": True,
                "role": CustomUser.Role.USER,
            },
        )

        validate_login_allowed(user)

        update_fields = []

        if not user.first_name and first_name:
            user.first_name = first_name
            update_fields.append("first_name")

        if not user.last_name and last_name:
            user.last_name = last_name
            update_fields.append("last_name")

        if not user.is_verified:
            user.is_verified = True
            update_fields.append("is_verified")

        update_fields.extend(_persist_remote_profile_image(user, picture))

        if update_fields:
            update_fields.append("updated_at")
            user.save(update_fields=list(dict.fromkeys(update_fields)))

        return Response(
            _finalize_device_login(
                user,
                request,
                is_new=is_new,
                include_is_new=True,
            ),
            status=status.HTTP_200_OK,
        )


# ============================================================
# USER DASHBOARD
# ============================================================

from gaming.models import SpinnerReward
from gaming.serializers import UserSpinnerRewardSerializer


class UserDashboardAPIView(DeviceActivityMixin, APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user
        today = timezone.localdate()

        games = GamingItem.objects.filter(
            is_active=True,
            is_deleted=False,
            maintenance_mode=False,
        ).select_related("category")

        recent_bookings = (
            Booking.objects.filter(
                user=user,
                is_deleted=False,
            )
            .select_related("item")
            .prefetch_related("members")
            .order_by("-created_at")[:5]
        )

        event_bookings = (
            EventBooking.objects.filter(
                user=user,
                is_deleted=False,
            )
            .select_related("event")
            .order_by("-created_at")
        )

        exclusive_events = ExclusiveEvent.objects.filter(
            event_date__gte=today,
            is_active=True,
            is_deleted=False,
        ).order_by("event_date", "start_time")[:6]

        combo_packs = ComboPack.objects.filter(
            is_active=True,
            is_deleted=False,
        )

        loyalty_transactions = LoyaltyTransaction.objects.filter(
            user=user,
            is_deleted=False,
        ).order_by("-created_at")[:10]

        spinner_history = (
            SpinnerSpin.objects.filter(
                user=user,
                is_deleted=False,
            )
            .select_related("reward")
            .order_by("-spun_at")[:5]
        )

        today = timezone.localdate()

        current_week = today.isocalendar()[1]
        current_year = today.year

        has_spun_this_week = SpinnerSpin.objects.filter(
            user=user,
            spin_week=current_week,
            spin_year=current_year,
            is_deleted=False,
        ).exists()
        spinner_rewards = SpinnerReward.objects.filter(
            is_active=True,
            is_deleted=False,
        )

        upcoming_booking = (
            Booking.objects.filter(
                user=user,
                is_deleted=False,
                booking_date__gte=today,
                status__in=[
                    Booking.Status.CONFIRMED,
                    Booking.Status.PENDING,
                ],
            )
            .select_related("item")
            .order_by("booking_date", "start_time")
            .first()
        )

        return Response(
            {
                "user_details": UserDetailsSerializer(
                    user,
                    context={"request": request},
                ).data,
                "dashboard_stats": {
                    "total_bookings": Booking.objects.filter(
                        user=user,
                        is_deleted=False,
                    ).count(),
                    "total_events": EventBooking.objects.filter(
                        user=user,
                        is_deleted=False,
                    ).count(),
                    "total_spins": SpinnerSpin.objects.filter(
                        user=user,
                        is_deleted=False,
                    ).count(),
                    "loyalty_points": user.loyalty_points,
                },
                "games_list": GameSerializer(
                    games,
                    many=True,
                    context={"request": request},
                ).data,
                "spinner": {
                    "can_spin": not has_spun_this_week,
                    "spin_history": SpinnerSpinSerializer(
                        spinner_history,
                        many=True,
                    ).data,
                },
                "upcoming_booking": (
                    BookingSerializer(
                        upcoming_booking,
                        context={"request": request},
                    ).data
                    if upcoming_booking
                    else None
                ),
                "loyalty_points": {
                    "current_points": user.loyalty_points,
                    "transactions": LoyaltyTransactionSerializer(
                        loyalty_transactions,
                        many=True,
                    ).data,
                },
                "recent_bookings": BookingSerializer(
                    recent_bookings,
                    many=True,
                    context={"request": request},
                ).data,
                "event_bookings": EventBookingSerializer(
                    event_bookings,
                    many=True,
                    context={"request": request},
                ).data,
                "exclusive_events": ExclusiveEventSerializer(
                    exclusive_events,
                    many=True,
                    context={"request": request},
                ).data,
                "combo_offers": ComboPackSerializer(
                    combo_packs,
                    many=True,
                    context={"request": request},
                ).data,
                "spinner": {
                    "can_spin": not has_spun_this_week,
                    "rewards": UserSpinnerRewardSerializer(
                        spinner_rewards, many=True
                    ).data,
                    "spin_history": SpinnerSpinSerializer(
                        spinner_history,
                        many=True,
                    ).data,
                },
            },
            status=status.HTTP_200_OK,
        )


import random
from django.db import IntegrityError


class UserSpinView(DeviceActivityMixin, APIView):
    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        user = request.user
        today = timezone.localdate()
        current_week = today.isocalendar()[1]
        current_year = today.year

        # Guard: one spin per week
        if SpinnerSpin.objects.filter(
            user=user,
            spin_week=current_week,
            spin_year=current_year,
            is_deleted=False,
        ).exists():
            return Response(
                {"detail": "You have already used your spin this week."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Weighted random reward selection
        rewards = SpinnerReward.objects.filter(is_active=True, is_deleted=False)
        if not rewards.exists():
            return Response(
                {"detail": "No rewards configured."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        population = list(rewards)
        weights = [r.probability for r in population]
        reward = random.choices(population, weights=weights, k=1)[0]

        # Save spin — unique_together is a safety net against race conditions
        try:
            spin = SpinnerSpin.objects.create(user=user, reward=reward)
        except IntegrityError:
            return Response(
                {"detail": "You have already used your spin this week."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Award loyalty points if it's a points reward
        if (
            reward.reward_type == SpinnerReward.RewardType.POINTS
            and reward.reward_points > 0
        ):
            user.loyalty_points += reward.reward_points
            user.save(update_fields=["loyalty_points"])
            LoyaltyTransaction.objects.create(
                user=user,
                points=reward.reward_points,
                transaction_type=LoyaltyTransaction.TransactionType.SPINNER,
                description=f"Spinner win: {reward.reward_name}",
            )

        return Response(
            {
                "reward": SpinnerRewardSerializer(reward).data,
                "spin_id": spin.id,
            },
            status=status.HTTP_200_OK,
        )


# ============================================================
# NAVBAR PROFILE
# ============================================================


class NavbarProfileView(DeviceActivityMixin, APIView):
    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user

        if not user.is_active:
            return Response(
                {"detail": "This account has been deactivated."},
                status=status.HTTP_403_FORBIDDEN,
            )

        if user.is_deleted:
            return Response(
                {"detail": "This account is no longer available."},
                status=status.HTTP_403_FORBIDDEN,
            )

        serializer = NavbarUserSerializer(
            user,
            context={"request": request},
        )

        return Response(
            {"user": serializer.data},
            status=status.HTTP_200_OK,
        )


# ============================================================
# ADMIN DASHBOARD
# ============================================================
from django.db.models import Exists, OuterRef


class AdminDashboardAPIView(DeviceActivityMixin, APIView):
    permission_classes = [
        permissions.IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):
        today = timezone.localdate()

        users = CustomUser.objects.filter(
            is_deleted=False,
            role=CustomUser.Role.USER,
        )

        bookings = Booking.objects.filter(
            is_deleted=False,
        )

        gaming_items = GamingItem.objects.filter(
            is_deleted=False,
        )

        total_revenue = (
            bookings.filter(
                payment_status=Booking.PaymentStatus.PAID,
            )
            .aggregate(total=Sum("total_amount"))
            .get("total")
            or 0
        )

        recent_bookings = bookings.select_related("user", "item").order_by(
            "-created_at"
        )[:5]

        recent_users = users.annotate(
            is_online_flag=Exists(
                UserDevice.objects.filter(
                    user=OuterRef("pk"),
                    is_deleted=False,
                    is_online=True,
                )
            )
        ).order_by("-created_at")[:5]

        attention_items = (
            gaming_items.filter(Q(maintenance_mode=True) | Q(is_active=False))
            .select_related("category")
            .order_by("-updated_at")[:10]
        )
        from django.db.models import F

        devices = (
            UserDevice.objects.filter(is_deleted=False)
            .select_related("user")
            .annotate(is_online_flag=F("is_online"))
            .order_by("-last_seen")[:20]
        )

        data = {
            "stats": {
                "total_users": users.count(),
                "total_bookings": bookings.count(),
                "total_revenue": total_revenue,
                "active_games": gaming_items.filter(
                    is_active=True,
                    maintenance_mode=False,
                ).count(),
                "pending_bookings": bookings.filter(
                    status=Booking.Status.PENDING,
                ).count(),
                "upcoming_events": ExclusiveEvent.objects.filter(
                    is_deleted=False,
                    is_active=True,
                    event_date__gte=today,
                ).count(),
            },
            "recent_bookings": AdminDashboardBookingSerializer(
                recent_bookings,
                many=True,
                context={"request": request},
            ).data,
            "recent_users": AdminDashboardUserSerializer(
                recent_users,
                many=True,
                context={"request": request},
            ).data,
            "low_stock_or_maintenance": AdminDashboardGamingItemSerializer(
                attention_items,
                many=True,
                context={"request": request},
            ).data,
            "recent_devices": AdminDashboardDeviceSerializer(
                devices,
                many=True,
            ).data,
        }

        return Response(
            data,
            status=status.HTTP_200_OK,
        )


from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from gaming.models import (
    GameCategory,
    GamingItem,
    ComboPack,
    ExclusiveEvent,
)

from gaming.serializers import (
    AdminGameCategorySerializer,
    AdminGamingItemSerializer,
    AdminComboPackSerializer,
    AdminExclusiveEventSerializer,
)


def admin_only(user):
    return user.is_authenticated and user.is_superuser


class AdminCategoryListCreateView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        if not admin_only(request.user):
            return Response(status=403)

        serializer = AdminGameCategorySerializer(
            GameCategory.objects.all(),
            many=True,
            context={"request": request},
        )

        return Response(serializer.data)

    def post(self, request):

        if not admin_only(request.user):
            return Response(status=403)

        serializer = AdminGameCategorySerializer(
            data=request.data,
            context={"request": request},
        )

        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data, status=status.HTTP_201_CREATED)


class AdminGameListCreateView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        serializer = AdminGamingItemSerializer(
            GamingItem.objects.select_related("category"),
            many=True,
            context={"request": request},
        )

        return Response(serializer.data)

    def post(self, request):

        serializer = AdminGamingItemSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data, status=201)


class AdminComboPackListCreateView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        serializer = AdminComboPackSerializer(
            ComboPack.objects.prefetch_related("gaming_items"),
            many=True,
            context={"request": request},
        )

        return Response(serializer.data)

    def post(self, request):

        serializer = AdminComboPackSerializer(
            data=request.data,
            context={"request": request},
        )

        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data, status=201)


class AdminEventListCreateView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        serializer = AdminExclusiveEventSerializer(
            ExclusiveEvent.objects.all(),
            many=True,
            context={"request": request},
        )

        return Response(serializer.data)

    def post(self, request):

        serializer = AdminExclusiveEventSerializer(
            data=request.data,
            context={"request": request},
        )

        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data, status=201)


class AdminGameDetailView(APIView):

    permission_classes = [IsAuthenticated]

    def put(self, request, pk):

        game = get_object_or_404(GamingItem, pk=pk)

        serializer = AdminGamingItemSerializer(
            game,
            data=request.data,
            partial=True,
            context={"request": request},
        )

        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data)

    def delete(self, request, pk):

        GamingItem.objects.filter(pk=pk).delete()

        return Response(status=204)


class AdminCategoryDetailView(APIView):

    permission_classes = [IsAuthenticated]

    def put(self, request, pk):

        category = GameCategory.objects.get(pk=pk)

        serializer = AdminGameCategorySerializer(
            category,
            data=request.data,
            partial=True,
            context={"request": request},
        )

        serializer.is_valid(raise_exception=True)
        serializer.save()

        return Response(serializer.data)

    def delete(self, request, pk):

        GameCategory.objects.filter(pk=pk).delete()

        return Response(status=204)


class AdminComboPackDetailView(APIView):

    permission_classes = [IsAuthenticated]

    def put(self, request, pk):

        combo = get_object_or_404(ComboPack, pk=pk)

        serializer = AdminComboPackSerializer(
            combo,
            data=request.data,
            partial=True,
            context={"request": request},
        )

        serializer.is_valid(raise_exception=True)

        serializer.save()

        return Response(serializer.data)

    def delete(self, request, pk):

        combo = get_object_or_404(ComboPack, pk=pk)

        combo.delete()

        return Response(status=status.HTTP_204_NO_CONTENT)


class AdminEventDetailView(APIView):

    permission_classes = [IsAuthenticated]

    def put(self, request, pk):

        event = get_object_or_404(ExclusiveEvent, pk=pk)

        serializer = AdminExclusiveEventSerializer(
            event,
            data=request.data,
            partial=True,
            context={"request": request},
        )

        serializer.is_valid(raise_exception=True)

        serializer.save()

        return Response(serializer.data)

    def delete(self, request, pk):

        event = get_object_or_404(ExclusiveEvent, pk=pk)

        event.delete()

        return Response(status=status.HTTP_204_NO_CONTENT)


from django.db.models import Q


class AdminUserListView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        if not request.user.is_superuser:
            return Response({"detail": "Permission denied"}, status=403)

        search = request.GET.get("search")

        queryset = CustomUser.objects.filter(role="user")

        if search:
            queryset = queryset.filter(
                Q(first_name__icontains=search)
                | Q(last_name__icontains=search)
                | Q(email__icontains=search)
                | Q(phone__icontains=search)
            )

        serializer = AdminUserListSerializer(
            queryset.order_by("-created_at"),
            many=True,
            context={"request": request},
        )

        return Response(serializer.data)


class AdminUserDetailView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request, pk):

        if not request.user.is_superuser:
            return Response({"detail": "Permission denied"}, status=403)

        user = get_object_or_404(CustomUser, pk=pk)

        serializer = AdminUserDetailSerializer(
            user,
            context={"request": request},
        )

        return Response(serializer.data)


class AdminUserStatusView(APIView):

    permission_classes = [IsAuthenticated]

    def patch(self, request, pk):

        if not request.user.is_superuser:
            return Response({"detail": "Permission denied"}, status=403)

        user = get_object_or_404(CustomUser, pk=pk)

        user.is_active = request.data.get("is_active", True)

        user.save()

        return Response({"success": True})


class AdminUserDeleteView(APIView):

    permission_classes = [IsAuthenticated]

    def delete(self, request, pk):

        if not request.user.is_superuser:
            return Response({"detail": "Permission denied"}, status=403)

        user = get_object_or_404(CustomUser, pk=pk)

        user.delete()

        return Response(status=204)


class AdminEventBookingListCreateView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        queryset = EventBooking.objects.select_related(
            "event",
            "user",
        ).order_by("-created_at")

        serializer = AdminEventBookingSerializer(
            queryset, many=True, context={"request": request}
        )

        return Response(serializer.data)

    def post(self, request):

        serializer = AdminEventBookingCreateSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        booking = serializer.save()

        return Response(
            AdminEventBookingDetailSerializer(
                booking, context={"request": request}
            ).data,
            status=201,
        )


class AdminEventBookingDetailView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request, pk):

        booking = get_object_or_404(
            EventBooking,
            pk=pk,
        )

        return Response(
            AdminEventBookingDetailSerializer(
                booking, context={"request": request}
            ).data
        )


from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from django.shortcuts import get_object_or_404
from django.utils import timezone

from gaming.models import EventBooking
from gaming.permissions import IsAdminUserRole
from gaming.utils.qr import generate_qr_code
from gaming.utils.emailjs import send_emailjs


class AdminEventBookingApproveView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request, pk):

        booking = get_object_or_404(EventBooking, pk=pk)

        if booking.status == EventBooking.Status.APPROVED:
            return Response(
                {"detail": "Booking is already approved."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # ── 1. Update booking status ──────────────────────────
        booking.status = EventBooking.Status.APPROVED
        booking.approved_by = request.user
        booking.approved_at = timezone.now()

        # ── 2. Generate QR code (once) ────────────────────────
        if not booking.qr_code:
            qr_file = generate_qr_code(
                data=str(booking.qr_token),
                filename=f"event_qr_{booking.booking_id}.png",
            )
            _store_encrypted_qr_code(booking, qr_file)

        booking.save()

        # ── 3. Build decrypted QR data URI for email ────────────
        qr_url = safe_image_to_data_uri(booking.qr_code) or ""
        # ── 4. Resolve recipient email ────────────────────────
        recipient_email = booking.user.email if booking.user else booking.email
        recipient_name = (
            booking.user.full_name if booking.user else (booking.full_name or "Guest")
        )

        if not recipient_email:
            return Response(
                {
                    "success": True,
                    "booking_id": booking.booking_id,
                    "warning": "Approved but no email address found — mail not sent.",
                }
            )

        # ── 5. Send email via EmailJS ─────────────────────────
        event = booking.event

        template_params = {
            "to_email": recipient_email,
            "customer_name": recipient_name,
            "booking_id": booking.booking_id,
            "event_name": event.title,
            "event_date": event.event_date.strftime("%d %B %Y"),
            "event_location": getattr(event, "location", "DQD Gaming Arena"),
            "amount_paid": str(booking.amount_paid),
            "qr_code_url": qr_url,
        }

        mail_sent = send_emailjs(template_params)
        if mail_sent:
            booking.qr_sent = True
            booking.save(update_fields=["qr_sent"])

        return Response(
            {
                "success": True,
                "booking_id": booking.booking_id,
                "email": recipient_email,
                "qr_sent": booking.qr_sent,
                "mail_sent": mail_sent,
            }
        )


class AdminEventBookingRejectView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request, pk):

        booking = get_object_or_404(
            EventBooking,
            pk=pk,
        )

        booking.status = EventBooking.Status.REJECTED

        booking.save()

        return Response({"success": True})


import uuid

from django.shortcuts import get_object_or_404
from django.utils import timezone
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView


class AdminVerifyEventQRView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request):

        qr_token = (request.data.get("qr_token") or "").strip()

        if not qr_token:
            return Response(
                {"valid": False, "error": "QR token is required."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            uuid.UUID(qr_token)
        except (ValueError, AttributeError, TypeError):
            return Response(
                {"valid": False, "error": "Invalid QR code format."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        booking = EventBooking.objects.filter(qr_token=qr_token).first()
        if booking is None:
            return Response(
                {"valid": False, "error": "No booking found for this QR code."},
                status=status.HTTP_404_NOT_FOUND,
            )

        # Optional: stop a code from being "checked in" twice.
        # Drop this block if you want re-scans to just succeed again.
        if booking.checked_in:
            return Response(
                {
                    "valid": False,
                    "error": "This booking has already been checked in.",
                    "booking_id": booking.booking_id,
                    "name": booking.full_name,
                    "event": booking.event.title,
                },
                status=status.HTTP_409_CONFLICT,
            )

        booking.checked_in = True
        booking.checked_in_at = timezone.now()
        booking.status = EventBooking.Status.ATTENDED
        booking.save()

        return Response(
            {
                "valid": True,
                "booking_id": booking.booking_id,
                "name": booking.full_name,
                "event": booking.event.title,
            }
        )


from gaming.utils.loyalty import (
    award_booking_loyalty_points,
)
from gaming.utils.emailjs import send_game_booking_email
from django.db import transaction


# ===========================================================================
# AdminBookingListCreateView  — replace your existing post() entirely
# ===========================================================================
class AdminBookingListCreateView(APIView):

    permission_classes = [IsAuthenticated, IsAdminUserRole]

    def get(self, request):
        queryset = (
            Booking.objects.select_related("user", "item", "combo_pack")
            .prefetch_related("members")
            .order_by("-created_at")
        )
        serializer = AdminBookingListSerializer(
            queryset, many=True, context={"request": request}
        )
        return Response(serializer.data)

    @transaction.atomic
    def post(self, request):
        serializer = AdminBookingCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        booking = serializer.save()
        award_booking_loyalty_points(
            booking=booking,
            admin_user=request.user,
        )
        # Refresh so members / related objects are available
        booking.refresh_from_db()

        # Send confirmation email immediately on admin-create
        # (no QR yet at create time — QR is generated on approve)
        # If you want to skip email until approve, remove the next two lines.
        send_game_booking_email(booking, request)

        return Response(
            AdminBookingDetailSerializer(booking, context={"request": request}).data,
            status=201,
        )


class AdminBookingDetailView(APIView):

    permission_classes = [IsAuthenticated, IsAdminUserRole]

    def get(self, request, pk):
        booking = get_object_or_404(Booking, pk=pk)
        return Response(
            AdminBookingDetailSerializer(booking, context={"request": request}).data
        )

    from decimal import Decimal

    def patch(self, request, pk):

        booking = get_object_or_404(
            Booking,
            pk=pk,
        )

        old_payment_status = booking.payment_status

        allowed_fields = {
            "status",
            "payment_status",
            "notes",
        }

        data = {k: v for k, v in request.data.items() if k in allowed_fields}

        for field, value in data.items():
            setattr(booking, field, value)

        # Add payment only once
        if (
            old_payment_status != Booking.PaymentStatus.PAID
            and booking.payment_status == Booking.PaymentStatus.PAID
        ):
            booking.price_paid = (
                booking.price_paid or Decimal("0.00")
            ) + booking.total_amount

        booking.save()

        return Response(
            AdminBookingDetailSerializer(
                booking,
                context={"request": request},
            ).data
        )


# ===========================================================================
# AdminBookingApproveView  — replace your existing post() entirely
# ===========================================================================
class AdminBookingApproveView(APIView):

    permission_classes = [IsAuthenticated, IsAdminUserRole]

    def post(self, request, pk):
        booking = get_object_or_404(Booking, pk=pk)

        booking.status = Booking.Status.CONFIRMED
        booking.approved_by = request.user
        booking.approved_at = timezone.now()

        # Generate QR if not already present
        if not booking.qr_code:
            qr_file = generate_qr_code(
                data=str(booking.qr_token),
                filename=f"booking_{booking.booking_id}.png",
            )
            _store_encrypted_qr_code(booking, qr_file)

        booking.save()

        # Send template 2 email (game booking confirmation with QR)
        mail_sent = send_game_booking_email(booking, request)

        if mail_sent:
            booking.qr_sent = True
            booking.save(update_fields=["qr_sent"])

        return Response(
            {
                "success": True,
                "mail_sent": mail_sent,
                "booking_id": booking.booking_id,
            }
        )


class AdminBookingRejectView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request, pk):

        booking = get_object_or_404(
            Booking,
            pk=pk,
        )

        booking.status = Booking.Status.CANCELLED

        booking.save()

        return Response({"success": True})


class AdminBookingVerifyQRView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request):

        qr_token = request.data.get("qr_token")

        booking = Booking.objects.filter(qr_token=qr_token).first()

        if not booking:

            return Response(
                {
                    "valid": False,
                    "message": "Invalid QR",
                },
                status=400,
            )

        if booking.checked_in:

            return Response(
                {
                    "valid": False,
                    "message": "Already Checked In",
                },
                status=400,
            )

        if not booking.is_qr_valid:

            return Response(
                {
                    "valid": False,
                    "message": "QR Expired",
                },
                status=400,
            )

        booking.checked_in = True

        booking.checked_in_at = timezone.now()

        booking.save()

        return Response(
            {
                "valid": True,
                "booking_id": booking.booking_id,
                "customer": booking.customer_name,
                "item": booking.item.name,
            }
        )


class AdminGameBookingListView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        queryset = (
            Booking.objects.filter(combo_pack__isnull=True)
            .select_related(
                "user",
                "item",
            )
            .order_by("-created_at")
        )

        return Response(
            AdminBookingListSerializer(
                queryset,
                many=True,
            ).data
        )


class AdminComboBookingListView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        queryset = (
            Booking.objects.filter(combo_pack__isnull=False)
            .select_related(
                "user",
                "combo_pack",
                "item",
            )
            .order_by("-created_at")
        )

        return Response(
            AdminBookingListSerializer(
                queryset,
                many=True,
            ).data
        )


class AdminHappyHourTemplateListCreateView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        slots = HappyHourTemplateSlot.objects.all()

        return Response(
            HappyHourTemplateSlotSerializer(
                slots,
                many=True,
            ).data
        )

    def post(self, request):

        serializer = HappyHourTemplateSlotSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        serializer.save()

        return Response(
            serializer.data,
            status=201,
        )


class AdminSpinnerRewardListCreateView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        rewards = SpinnerReward.objects.select_related("template_slot")

        return Response(
            SpinnerRewardSerializer(
                rewards,
                many=True,
            ).data
        )

    def post(self, request):

        serializer = SpinnerRewardSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        serializer.save()

        return Response(
            serializer.data,
            status=201,
        )


class AdminHappyHourBookingsView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        bookings = HappyHourBooking.objects.select_related(
            "happy_hour_slot",
            "booking",
        ).order_by("-created_at")

        return Response(
            HappyHourBookingSerializer(
                bookings,
                many=True,
            ).data
        )


class AdminSpinnerSpinsView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        spins = SpinnerSpin.objects.select_related(
            "user",
            "reward",
        ).order_by("-spun_at")

        data = []

        for spin in spins:
            data.append(
                {
                    "id": spin.id,
                    "user": spin.user.full_name,
                    "reward": spin.reward.reward_name,
                    "week": spin.spin_week,
                    "year": spin.spin_year,
                    "spun_at": spin.spun_at,
                }
            )

        return Response(data)


class AdminAssignHappyHourGameView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request, pk):

        slot = get_object_or_404(
            HappyHourSlot,
            pk=pk,
        )

        serializer = AdminAssignHappyHourGameSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        game = get_object_or_404(
            GamingItem, pk=serializer.validated_data["gaming_item"]
        )

        slot.gaming_item = game
        slot.save(update_fields=["gaming_item"])

        return Response({"success": True, "message": "Game assigned successfully."})


class AdminVerifyHappyHourView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request):

        slot_id = request.data.get("slot_id")

        slot = get_object_or_404(
            HappyHourSlot,
            pk=slot_id,
        )

        if slot.is_used:

            return Response(
                {
                    "valid": False,
                    "message": "Already redeemed",
                }
            )

        slot.is_used = True
        slot.redeemed_at = timezone.now()

        slot.save()

        return Response(
            {
                "valid": True,
                "user": slot.user.full_name,
                "slot": slot.template_slot.slot_name,
            }
        )


class AdminUserLoyaltyListView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        users = CustomUser.objects.filter(role="user").order_by("-created_at")

        serializer = AdminUserLoyaltySerializer(
            users,
            many=True,
        )

        return Response(serializer.data)


class AdminUserLoyaltyHistoryView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request, pk):

        transactions = (
            LoyaltyTransaction.objects.filter(user_id=pk)
            .select_related(
                "granted_by",
                "booking",
            )
            .order_by("-created_at")
        )

        return Response(
            LoyaltyTransactionSerializer(
                transactions,
                many=True,
            ).data
        )


class AdminAdjustLoyaltyView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request, pk):

        serializer = AdminAdjustLoyaltySerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        user = get_object_or_404(
            CustomUser,
            pk=pk,
        )

        points = serializer.validated_data["points"]

        action = serializer.validated_data["action"]

        reason = serializer.validated_data["reason"]

        if action == "add":

            user.loyalty_points += points

            LoyaltyTransaction.objects.create(
                user=user,
                points=points,
                transaction_type="admin_grant",
                granted_by=request.user,
                description=reason,
            )

        else:

            if user.loyalty_points < points:

                return Response(
                    {"message": "Insufficient points."},
                    status=400,
                )

            user.loyalty_points -= points

            LoyaltyTransaction.objects.create(
                user=user,
                points=-points,
                transaction_type="redeemed",
                granted_by=request.user,
                description=reason,
            )

        user.save(update_fields=["loyalty_points"])

        return Response(
            {
                "success": True,
                "current_points": user.loyalty_points,
            }
        )


from datetime import timedelta
import openpyxl

from django.http import HttpResponse

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
)

from reportlab.lib import colors
from django.db.models import (
    Sum,
    Avg,
    Count,
    Q,
)

from django.db.models.functions import (
    TruncDate,
    TruncWeek,
    TruncMonth,
)

from django.utils import timezone

from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
import gspread

from google.oauth2.service_account import Credentials

from django.conf import settings

import os
from gaming.models import Booking
from gaming.permissions import IsAdminUserRole


def get_booking_queryset(request):

    queryset = Booking.objects.filter(is_deleted=False)

    period = request.GET.get(
        "period",
        "monthly",
    )

    today = timezone.localdate()

    if period == "daily":

        queryset = queryset.filter(booking_date=today)

    elif period == "weekly":

        week_start = today - timedelta(days=today.weekday())

        queryset = queryset.filter(booking_date__gte=week_start)

    elif period == "monthly":

        queryset = queryset.filter(
            booking_date__month=today.month,
            booking_date__year=today.year,
        )

    elif period == "yearly":

        queryset = queryset.filter(
            booking_date__year=today.year,
        )

    elif period == "custom":

        start = request.GET.get("start_date")

        end = request.GET.get("end_date")

        if start and end:

            queryset = queryset.filter(
                booking_date__range=[
                    start,
                    end,
                ]
            )

    return queryset


class AdminAccountingDashboardAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        bookings = get_booking_queryset(request)

        completed = bookings.filter(status=Booking.Status.COMPLETED)

        total_bookings = bookings.count()

        completed_bookings = completed.count()

        cancelled = bookings.filter(status=Booking.Status.CANCELLED).count()

        pending = bookings.filter(status=Booking.Status.PENDING).count()

        gross = completed.aggregate(total=Sum("subtotal"))["total"] or 0

        discount = completed.aggregate(total=Sum("discount_amount"))["total"] or 0

        revenue = completed.aggregate(total=Sum("total_amount"))["total"] or 0

        received = completed.aggregate(total=Sum("price_paid"))["total"] or 0

        average = completed.aggregate(avg=Avg("total_amount"))["avg"] or 0

        pending_payment = (
            bookings.filter(payment_status=Booking.PaymentStatus.PENDING).aggregate(
                total=Sum("total_amount")
            )["total"]
            or 0
        )

        return Response(
            {
                "summary": {
                    "total_bookings": total_bookings,
                    "completed_bookings": completed_bookings,
                    "cancelled_bookings": cancelled,
                    "pending_bookings": pending,
                    "gross_revenue": gross,
                    "discount": discount,
                    "net_revenue": revenue,
                    "amount_received": received,
                    "pending_payment": pending_payment,
                    "average_booking": average,
                    "profit": received,
                    "loss": discount,
                }
            }
        )


from django.db.models.functions import (
    TruncDate,
    TruncWeek,
    TruncMonth,
    ExtractYear,
)


class AdminRevenueChartAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        period = request.GET.get("period", "monthly")

        queryset = get_booking_queryset(request).filter(status=Booking.Status.COMPLETED)

        if period == "daily":
            queryset = queryset.annotate(label=TruncDate("booking_date"))

        elif period == "weekly":
            queryset = queryset.annotate(label=TruncWeek("booking_date"))

        elif period == "monthly":
            queryset = queryset.annotate(label=TruncMonth("booking_date"))

        elif period == "yearly":
            queryset = queryset.annotate(label=ExtractYear("booking_date"))

        else:  # custom
            queryset = queryset.annotate(label=TruncDate("booking_date"))

        data = (
            queryset.values("label")
            .annotate(
                bookings=Count("id"),
                revenue=Sum("price_paid"),
                discount=Sum("discount_amount"),
            )
            .order_by("label")
        )

        results = []

        for row in data:
            results.append(
                {
                    "label": str(row["label"]),
                    "bookings": row["bookings"] or 0,
                    "revenue": float(row["revenue"] or 0),
                    "discount": float(row["discount"] or 0),
                }
            )

        return Response(results)


class AdminAccountingBookingsAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        bookings = (
            get_booking_queryset(request)
            .select_related(
                "user",
                "item",
            )
            .order_by("-booking_date")
        )

        data = []

        for booking in bookings:

            data.append(
                {
                    "booking_id": booking.booking_id,
                    "customer": (
                        booking.customer_name or booking.user.full_name
                        if booking.user
                        else "-"
                    ),
                    "booking_date": booking.created_at.date(),
                    "game": (booking.item.name if booking.item else "Deleted Game"),
                    "hours": booking.total_hours or 0,
                    "subtotal": booking.subtotal or 0,
                    "discount": booking.discount_amount or 0,
                    "total": booking.total_amount or 0,
                    "paid": booking.price_paid or 0,
                    "payment_status": booking.payment_status,
                    "booking_status": booking.status,
                }
            )

        return Response(data)


from django.db.models import Sum
from rest_framework.views import APIView
from rest_framework.response import Response


class AdminProfitLossAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        bookings = get_booking_queryset(request)

        completed = bookings.filter(status=Booking.Status.COMPLETED)

        gross_income = completed.aggregate(total=Sum("subtotal"))["total"] or 0

        discount = completed.aggregate(total=Sum("discount_amount"))["total"] or 0

        revenue = completed.aggregate(total=Sum("price_paid"))["total"] or 0

        # Replace later with Expense model
        expenses = 0

        profit = revenue - expenses

        return Response(
            {
                "gross_income": gross_income,
                "discount": discount,
                "expenses": expenses,
                "net_revenue": revenue,
                "profit": profit,
                "loss": max(0, expenses - revenue),
            }
        )


import openpyxl

from django.http import HttpResponse


class AdminAccountingExcelExportAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        workbook = openpyxl.Workbook()

        sheet = workbook.active

        sheet.title = "Accounting"

        headers = [
            "Booking ID",
            "Customer",
            "Game",
            "Date",
            "Start Time",
            "End Time",
            "Hours",
            "Subtotal",
            "Discount",
            "Total",
            "Payment Status",
            "Paid",
            "Notes",
            "Status",
        ]

        for i, h in enumerate(headers, start=1):
            sheet.cell(row=1, column=i).value = h

        bookings = get_booking_queryset(request)

        row = 2

        for booking in bookings:

            sheet.cell(row=row, column=1).value = booking.booking_id
            sheet.cell(row=row, column=2).value = booking.customer_name
            sheet.cell(row=row, column=3).value = (
                booking.item.name if booking.item else "-"
            )
            sheet.cell(row=row, column=4).value = str(booking.booking_date)
            sheet.cell(row=row, column=5).value = str(booking.start_time)
            sheet.cell(row=row, column=6).value = str(booking.end_time)

            sheet.cell(row=row, column=7).value = float(booking.total_hours)
            sheet.cell(row=row, column=8).value = float(booking.subtotal)
            sheet.cell(row=row, column=9).value = float(booking.discount_amount)
            sheet.cell(row=row, column=10).value = float(booking.total_amount)

            sheet.cell(row=row, column=11).value = str(booking.payment_status)
            sheet.cell(row=row, column=12).value = float(booking.price_paid)
            sheet.cell(row=row, column=13).value = str(booking.notes)

            sheet.cell(row=row, column=14).value = booking.status

            row += 1

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        response["Content-Disposition"] = 'attachment; filename="Accounting.xlsx"'

        workbook.save(response)

        return response


from django.http import HttpResponse

from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
)

from reportlab.lib import colors
from decimal import Decimal


from decimal import Decimal
from datetime import date, datetime, time

from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle


def sheet_value(value):
    if value is None:
        return ""

    if isinstance(value, Decimal):
        return float(value)

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, time):
        return value.strftime("%H:%M:%S")

    return value


class AdminAccountingPDFExportAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get(self, request):

        response = HttpResponse(content_type="application/pdf")

        response["Content-Disposition"] = 'attachment; filename="Accounting.pdf"'

        doc = SimpleDocTemplate(
            response,
            pagesize=landscape(A4),
            leftMargin=10,
            rightMargin=10,
            topMargin=15,
            bottomMargin=15,
        )

        data = [
            [
                "Booking ID",
                "Customer",
                "Game / Item",
                "Booking Date",
                "Start Time",
                "End Time",
                "Total Hours",
                "Subtotal",
                "Discount",
                "Total",
                "Payment",
                "Paid",
                "Notes",
                "Status",
            ]
        ]

        bookings = get_booking_queryset(request)

        for booking in bookings:

            data.append(
                [
                    sheet_value(booking.booking_id),
                    sheet_value(booking.customer_name),
                    sheet_value(
                        booking.item.name
                        if booking.item
                        else booking.combo_pack.name if booking.combo_pack else ""
                    ),
                    sheet_value(booking.booking_date),
                    sheet_value(booking.start_time),
                    sheet_value(booking.end_time),
                    sheet_value(booking.total_hours),
                    sheet_value(booking.subtotal),
                    sheet_value(booking.discount_amount),
                    sheet_value(booking.total_amount),
                    sheet_value(booking.payment_status),
                    sheet_value(booking.price_paid),
                    sheet_value(booking.notes),
                    sheet_value(booking.status),
                ]
            )

        table = Table(
            data,
            repeatRows=1,
            colWidths=[
                55,  # Booking ID
                75,  # Customer
                70,  # Game
                55,  # Date
                40,  # Start
                40,  # End
                50,  # Hours
                50,  # Subtotal
                50,  # Discount
                50,  # Total
                50,  # Payment
                50,  # Paid
                120,  # Notes
                45,  # Status
            ],
        )

        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.darkblue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ]
            )
        )

        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, Spacer

        styles = getSampleStyleSheet()

        elements = [
            Paragraph("<b>DQD Gaming Accounting Report</b>", styles["Title"]),
            Spacer(1, 12),
            table,
        ]

        doc.build(elements)

        return response


class AdminAccountingGoogleSheetSyncAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def post(self, request):

        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]

        credentials = Credentials.from_service_account_file(
            settings.GOOGLE_SERVICE_ACCOUNT_INFO,
            scopes=[
                "https://www.googleapis.com/auth/spreadsheets",
            ],
        )

        client = gspread.authorize(credentials)

        sheet = client.open_by_key(settings.GOOGLE_SHEET_ID).sheet1

        sheet.clear()

        rows = [
            [
                "Booking ID",
                "Customer",
                "Game / Item",
                "Booking Date",
                "Start Time",
                "End Time",
                "Total Hours",
                "Subtotal",
                "Discount",
                "Total",
                "Payment",
                "Paid",
                "Notes",
                "Status",
            ]
        ]

        bookings = Booking.objects.filter(is_deleted=False)

        for booking in bookings:

            rows.append(
                [
                    sheet_value(booking.booking_id),
                    sheet_value(booking.customer_name),
                    sheet_value(
                        booking.item.name
                        if booking.item
                        else booking.combo_pack.name if booking.combo_pack else ""
                    ),
                    sheet_value(booking.booking_date),
                    sheet_value(booking.start_time),
                    sheet_value(booking.end_time),
                    sheet_value(booking.total_hours),
                    sheet_value(booking.subtotal),
                    sheet_value(booking.discount_amount),
                    sheet_value(booking.total_amount),
                    sheet_value(booking.payment_status),
                    sheet_value(booking.price_paid),
                    sheet_value(booking.notes),
                    sheet_value(booking.status),
                ]
            )

        sheet.clear()

        sheet.update(
            values=rows,
            range_name="A1",
        )

        return Response(
            {
                "success": True,
                "message": "Google Sheet updated successfully.",
                "rows": len(rows) - 1,
            }
        )


from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status

from gaming.serializers import (
    ProfileSerializer,
    ChangePasswordSerializer,
)


class ProfileAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
    ]

    def get(self, request):

        serializer = ProfileSerializer(
            request.user,
            context={
                "request": request,
            },
        )

        return Response(serializer.data)

    def patch(self, request):

        serializer = ProfileSerializer(
            request.user,
            data=request.data,
            partial=True,
            context={
                "request": request,
            },
        )

        serializer.is_valid(raise_exception=True)

        user = serializer.save()

        return Response(
            ProfileSerializer(
                user,
                context={"request": request},
            ).data
        )


class ChangePasswordAPIView(APIView):

    permission_classes = [
        IsAuthenticated,
    ]

    def post(self, request):

        serializer = ChangePasswordSerializer(
            data=request.data,
            context={
                "request": request,
            },
        )

        serializer.is_valid(raise_exception=True)

        user = request.user

        user.set_password(serializer.validated_data["new_password"])

        user.save()

        return Response(
            {
                "success": True,
                "message": "Password changed successfully.",
            }
        )


from datetime import timedelta

from django.utils import timezone

from gaming.models import (
    CustomUser,
    PasswordResetOTP,
)

from gaming.serializers import (
    ForgotPasswordSerializer,
    VerifyForgotOTPSerializer,
    ResetPasswordSerializer,
)


class ForgotPasswordAPIView(APIView):

    permission_classes = []

    def post(self, request):

        serializer = ForgotPasswordSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        user = CustomUser.objects.get(email=serializer.validated_data["email"])

        PasswordResetOTP.objects.filter(
            user=user,
            is_used=False,
        ).delete()

        otp = PasswordResetOTP.generate_otp()

        PasswordResetOTP.objects.create(
            user=user,
            otp=otp,
            expires_at=timezone.now() + timedelta(minutes=10),
        )

        return Response(
            {
                "success": True,
                "message": "OTP generated successfully.",
                "otp": otp,
                "email": user.email,
            }
        )


class VerifyForgotOTPAPIView(APIView):

    permission_classes = []

    def post(self, request):

        serializer = VerifyForgotOTPSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        user = get_object_or_404(
            CustomUser,
            email=serializer.validated_data["email"],
        )

        otp = PasswordResetOTP.objects.filter(
            user=user,
            otp=serializer.validated_data["otp"],
            is_used=False,
        ).first()

        if not otp:

            return Response({"success": False, "message": "Invalid OTP."}, status=400)

        if otp.is_expired():

            otp.delete()

            return Response({"success": False, "message": "OTP expired."}, status=400)

        return Response({"success": True, "message": "OTP verified successfully."})


class ResetPasswordAPIView(APIView):

    permission_classes = []

    def post(self, request):

        serializer = ResetPasswordSerializer(data=request.data)

        serializer.is_valid(raise_exception=True)

        user = get_object_or_404(
            CustomUser,
            email=serializer.validated_data["email"],
        )

        otp = PasswordResetOTP.objects.filter(
            user=user,
            otp=serializer.validated_data["otp"],
            is_used=False,
        ).first()

        if not otp:

            return Response({"success": False, "message": "Invalid OTP."}, status=400)

        if otp.is_expired():

            otp.delete()

            return Response({"success": False, "message": "OTP expired."}, status=400)

        user.set_password(serializer.validated_data["new_password"])

        user.save()

        otp.is_used = True
        otp.save()

        return Response({"success": True, "message": "Password reset successfully."})


# ------------------------------ Home -------------------------------------


from rest_framework.views import APIView
from rest_framework.response import Response

from gaming.models import (
    GameCategory,
    GamingItem,
    ComboPack,
    ExclusiveEvent,
    CustomUser,
)

from gaming.serializers import (
    HomeSerializer,
)
from django.utils import timezone

now = timezone.localtime()

class PublicHomeAPIView(APIView):
    permission_classes = []
    throttle_classes = [ScopedRateThrottle]
    throttle_scope = "public_home"

    def get(self, request):

        # -----------------------
        # Events
        # -----------------------

        events = (
                ExclusiveEvent.objects.filter(
                    is_active=True,
                )
                .filter(
                    models.Q(event_date__gt=now.date()) |
                    models.Q(
                        event_date=now.date(),
                        end_time__gte=now.time(),
                    )
                )
                .order_by("event_date", "start_time")[:5]
            )

        # -----------------------
        # Combo Packs
        # -----------------------

        combos = ComboPack.objects.filter(is_active=True).prefetch_related(
            "gaming_items"
        )[:8]

        # -----------------------
        # Categories + Games
        # -----------------------

        categories = (
            GameCategory.objects.filter(is_active=True)
            .prefetch_related("gaming_items")
            .order_by("name")
        )

        # -----------------------
        # Featured Games
        # -----------------------

        featured_games = GamingItem.objects.filter(
            is_active=True,
            maintenance_mode=False,
        )[:5]

        # -----------------------
        # Banner Priority
        # -----------------------

        banners = []

        for event in events:

            banners.append(
                {
                    "id": event.id,
                    "title": event.title,
                    "subtitle": event.description,
                    "image": (
                        safe_image_to_data_uri(event.image)
                        if event.image
                        else None
                    ),
                    "type": "event",
                    "action_url": f"/events/{event.id}",
                }
            )

        for game in featured_games:

            banners.append(
                {
                    "id": game.id,
                    "title": game.name,
                    "subtitle": game.description,
                    "image": (
                        safe_image_to_data_uri(game.image)
                        if game.image
                        else None
                    ),
                    "type": "game",
                    "action_url": f"/games/{game.id}",
                }
            )

        # -----------------------
        # Statistics
        # -----------------------

        stats = {
            "total_users": CustomUser.objects.filter(
                role=CustomUser.Role.USER,
                is_active=True,
            ).count(),
            "total_games": GamingItem.objects.filter(is_active=True).count(),
            "total_events": ExclusiveEvent.objects.filter(is_active=True).count(),
            "total_combo_packs": ComboPack.objects.filter(is_active=True).count(),
        }

        serializer = HomeSerializer(
            {
                "banners": banners,
                "events": events,
                "combo_packs": combos,
                "categories": categories,
                "stats": stats,
            },
            context={
                "request": request,
            },
        )

        return Response(
            {
                "success": True,
                "data": serializer.data,
            }
        )


# --------------------------------- User --------------------------------

from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response


class UserGameCategoryAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        categories = GameCategory.objects.filter(is_active=True).order_by("name")

        serializer = UserGameCategorySerializer(
            categories, many=True, context={"request": request}
        )

        return Response(
            {"success": True, "count": categories.count(), "results": serializer.data}
        )


from django.db.models import Q


class UserGameListAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        category = request.GET.get("category")
        search = request.GET.get("search")

        games = (
            GamingItem.objects.select_related("category")
            .filter(is_active=True, maintenance_mode=False)
            .order_by("name")
        )

        if category:
            games = games.filter(category_id=category)

        if search:
            games = games.filter(
                Q(name__icontains=search) | Q(description__icontains=search)
            )

        serializer = UserGamingItemListSerializer(
            games, many=True, context={"request": request}
        )

        return Response(
            {"success": True, "count": games.count(), "results": serializer.data}
        )


from django.shortcuts import get_object_or_404


class UserGameDetailAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request, pk):

        game = get_object_or_404(
            GamingItem.objects.select_related("category"), pk=pk, is_active=True
        )

        serializer = UserGamingItemDetailSerializer(game, context={"request": request})

        return Response({"success": True, "data": serializer.data})


class UserBookingCreateAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def post(self, request):

        serializer = BookingCreateSerializer(
            data=request.data, context={"request": request}
        )

        serializer.is_valid(raise_exception=True)

        booking = serializer.save()

        return Response(
            {
                "success": True,
                "message": "Booking created successfully.",
                "booking": BookingDetailSerializer(
                    booking, context={"request": request}
                ).data,
            },
            status=201,
        )


from django.db.models import Q


class UserBookingListAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        bookings = (
            Booking.objects.filter(
                user=request.user,
                status__in=[
                    Booking.Status.PENDING,
                    Booking.Status.CONFIRMED,
                ],
            )
            .select_related(
                "item",
                "combo_pack",
            )
            .prefetch_related(
                "members",
            )
            .order_by(
                "booking_date",
                "start_time",
            )
        )

        serializer = BookingListSerializer(
            bookings,
            many=True,
            context={"request": request},
        )

        return Response(
            {
                "success": True,
                "count": bookings.count(),
                "results": serializer.data,
            }
        )


class UserBookingDetailAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request, pk):

        booking = get_object_or_404(
            Booking.objects.select_related(
                "item",
                "item__category",
                "combo_pack",
                "approved_by",
            ).prefetch_related(
                "members",
            ),
            pk=pk,
            user=request.user,
        )

        serializer = BookingDetailSerializer(
            booking,
            context={"request": request},
        )

        return Response(
            {
                "success": True,
                "data": serializer.data,
            }
        )


class UserBookingHistoryAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        bookings = (
            Booking.objects.filter(
                user=request.user,
                status__in=[
                    Booking.Status.COMPLETED,
                    Booking.Status.CANCELLED,
                ],
            )
            .select_related(
                "item",
                "combo_pack",
            )
            .order_by("-booking_date", "-start_time")
        )

        serializer = BookingHistorySerializer(
            bookings,
            many=True,
            context={"request": request},
        )

        return Response(
            {
                "success": True,
                "count": bookings.count(),
                "results": serializer.data,
            }
        )


class UserCancelBookingAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def put(self, request, pk):

        booking = get_object_or_404(
            Booking,
            pk=pk,
            user=request.user,
        )

        if booking.status != Booking.Status.PENDING:

            return Response(
                {
                    "success": False,
                    "message": "Only pending bookings can be cancelled.",
                },
                status=400,
            )

        booking.status = Booking.Status.CANCELLED
        booking.save(update_fields=["status"])

        return Response(
            {
                "success": True,
                "message": "Booking cancelled successfully.",
            }
        )


from datetime import datetime, timedelta, time


class UserAvailableSlotsAPIView(APIView):

    permission_classes = [IsAuthenticated]

    SLOT_DURATION = 60  # minutes

    def get(self, request, pk):

        date = request.GET.get("date")

        if not date:
            return Response(
                {"success": False, "message": "Date is required."},
                status=400,
            )

        booking_date = datetime.strptime(date, "%Y-%m-%d").date()

        item = get_object_or_404(
            GamingItem,
            pk=pk,
            is_active=True,
        )

        bookings = Booking.objects.filter(
            item=item,
            booking_date=booking_date,
        ).exclude(status=Booking.Status.CANCELLED)

        opening = time(9, 0)
        closing = time(23, 0)

        slots = []

        current = datetime.combine(
            booking_date,
            opening,
        )

        end = datetime.combine(
            booking_date,
            closing,
        )

        while current < end:

            slot_end = current + timedelta(minutes=self.SLOT_DURATION)

            available = True

            for booking in bookings:

                booking_start = datetime.combine(
                    booking_date,
                    booking.start_time,
                )

                booking_end = datetime.combine(
                    booking_date,
                    booking.end_time,
                )

                if current < booking_end and slot_end > booking_start:
                    available = False
                    break

            slots.append(
                {
                    "start_time": current.time(),
                    "end_time": slot_end.time(),
                    "available": available,
                }
            )

            current = slot_end

        serializer = AvailableSlotSerializer(
            slots,
            many=True,
        )

        return Response(
            {
                "success": True,
                "results": serializer.data,
            }
        )


from django.db.models import Prefetch
from django.utils import timezone
from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from gaming.models import Booking, BookingMember, EventBooking
from gaming.serializers import (
    CancelBookingSerializer,
    UserBookingWalletSerializer,
    UserEventBookingWalletSerializer,
)


class UserBookingWalletView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        now = timezone.now()
        today = timezone.localdate()

        bookings = (
            Booking.objects.filter(user=request.user, is_deleted=False)
            .select_related("item", "item__category", "combo_pack")
            .prefetch_related(Prefetch("members", queryset=BookingMember.objects.all()))
        )
        event_bookings = EventBooking.objects.filter(
            user=request.user, is_deleted=False
        ).select_related("event")

        active_statuses = [
            Booking.Status.PENDING,
            Booking.Status.CONFIRMED,
        ]
        active_event_statuses = [
            EventBooking.Status.PENDING,
            EventBooking.Status.APPROVED,
        ]

        upcoming_bookings = [
            booking
            for booking in bookings.filter(status__in=active_statuses)
            if timezone.make_aware(
                timezone.datetime.combine(booking.booking_date, booking.end_time)
            )
            >= now
        ]
        upcoming_events = [
            event_booking
            for event_booking in event_bookings.filter(status__in=active_event_statuses)
            if timezone.make_aware(
                timezone.datetime.combine(
                    event_booking.event.event_date, event_booking.event.end_time
                )
            )
            >= now
        ]

        history_bookings = [
            booking
            for booking in bookings.exclude(
                id__in=[item.id for item in upcoming_bookings]
            )
            if booking.booking_date <= today
            or booking.status in [Booking.Status.COMPLETED, Booking.Status.CANCELLED]
        ]
        history_events = [
            event_booking
            for event_booking in event_bookings.exclude(
                id__in=[item.id for item in upcoming_events]
            )
            if event_booking.event.event_date <= today
            or event_booking.status
            in [
                EventBooking.Status.ATTENDED,
                EventBooking.Status.CANCELLED,
                EventBooking.Status.REJECTED,
            ]
        ]

        upcoming = self._serialize_booking_group(
            request, upcoming_bookings, upcoming_events
        )
        history = self._serialize_booking_group(
            request, history_bookings, history_events
        )

        return Response(
            {
                "upcoming": sorted(upcoming, key=lambda item: item["starts_at"]),
                "history": sorted(
                    history, key=lambda item: item["starts_at"], reverse=True
                ),
                "counts": {
                    "upcoming": len(upcoming),
                    "history": len(history),
                    "combo": bookings.filter(combo_pack__isnull=False).count(),
                    "event": event_bookings.count(),
                },
            }
        )

    def _serialize_booking_group(self, request, bookings, event_bookings):
        context = {"request": request}
        booking_data = UserBookingWalletSerializer(
            bookings, many=True, context=context
        ).data
        event_data = UserEventBookingWalletSerializer(
            event_bookings, many=True, context=context
        ).data
        return [*booking_data, *event_data]


class UserCancelBookingView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, booking_id):
        serializer = CancelBookingSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)

        booking = (
            Booking.objects.filter(
                booking_id=booking_id, user=request.user, is_deleted=False
            )
            .select_related("item", "combo_pack")
            .first()
        )
        if booking:
            return self._cancel_regular_booking(request, booking, serializer)

        event_booking = (
            EventBooking.objects.filter(
                booking_id=booking_id, user=request.user, is_deleted=False
            )
            .select_related("event")
            .first()
        )
        if event_booking:
            return self._cancel_event_booking(request, event_booking)

        return Response(
            {"detail": "Booking not found."},
            status=status.HTTP_404_NOT_FOUND,
        )

    def _cancel_regular_booking(self, request, booking, serializer):
        if booking.status in [Booking.Status.CANCELLED, Booking.Status.COMPLETED]:
            return Response(
                {"detail": "This booking cannot be cancelled."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        booking_start = timezone.make_aware(
            timezone.datetime.combine(booking.booking_date, booking.start_time)
        )
        if timezone.now() >= booking_start:
            return Response(
                {"detail": "Bookings cannot be cancelled after they start."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        booking.status = Booking.Status.CANCELLED
        booking.cancelled_at = timezone.now()
        booking.cancelled_by = request.user
        booking.cancelled_from = "user"
        booking.cancellation_reason = serializer.validated_data.get("reason", "")
        if booking.payment_status == Booking.PaymentStatus.PAID:
            booking.refund_status = "pending"
            booking.refund_amount = booking.price_paid
        booking.save()

        return Response(
            UserBookingWalletSerializer(booking, context={"request": request}).data,
            status=status.HTTP_200_OK,
        )

    def _cancel_event_booking(self, request, event_booking):
        if event_booking.status in [
            EventBooking.Status.CANCELLED,
            EventBooking.Status.ATTENDED,
            EventBooking.Status.REJECTED,
        ]:
            return Response(
                {"detail": "This event booking cannot be cancelled."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        event_start = timezone.make_aware(
            timezone.datetime.combine(
                event_booking.event.event_date, event_booking.event.start_time
            )
        )
        if timezone.now() >= event_start:
            return Response(
                {"detail": "Event bookings cannot be cancelled after they start."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        event_booking.status = EventBooking.Status.CANCELLED
        event_booking.save(update_fields=["status", "updated_at"])

        return Response(
            UserEventBookingWalletSerializer(
                event_booking, context={"request": request}
            ).data,
            status=status.HTTP_200_OK,
        )















from rest_framework import status
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from gaming.serializers import (
    LoyaltyRedeemedBookingSerializer,
    LoyaltySlotRedemptionSerializer,
    UserGamingItemSerializer
)


class RedeemLoyaltySlotAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = LoyaltySlotRedemptionSerializer(
            data=request.data,
            context={"request": request},
        )
        serializer.is_valid(raise_exception=True)
        booking = serializer.save()

        return Response(
            {
                "message": "Loyalty points redeemed. Your game slot is booked!",
                "booking": LoyaltyRedeemedBookingSerializer(booking).data,
            },
            status=status.HTTP_201_CREATED,
        )


from datetime import datetime, timedelta, time
from django.shortcuts import get_object_or_404

class AvailableGamingSlotsAPIView(APIView):
    permission_classes = [IsAuthenticated]

    OPEN_TIME = time(10, 0)
    CLOSE_TIME = time(22, 0)

    def get(self, request, item_id):

        date = request.query_params.get("date")

        if not date:
            return Response(
                {"detail": "date is required"},
                status=400
            )

        item = get_object_or_404(
                GamingItem,
                pk=item_id,
                is_active=True,
                is_deleted=False,
            )

        bookings = Booking.objects.filter(
            item=item,
            booking_date=date
        ).exclude(
            status=Booking.Status.CANCELLED
        )

        current = datetime.combine(
            datetime.strptime(date, "%Y-%m-%d"),
            self.OPEN_TIME
        )

        end = datetime.combine(
            datetime.strptime(date, "%Y-%m-%d"),
            self.CLOSE_TIME
        )

        slots = []

        while current < end:

            slot_end = current + timedelta(hours=1)

            booked = bookings.filter(
                start_time__lt=slot_end.time(),
                end_time__gt=current.time()
            ).exists()

            slots.append({

                "start_time": current.time(),

                "end_time": slot_end.time(),

                "available": not booked

            })

            current += timedelta(hours=1)

        serializer = AvailableSlotSerializer(slots, many=True)
        return Response(serializer.data)


class UserLoyaltyAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response({
            "loyalty_points": request.user.loyalty_points,
            "redeemable_hours": request.user.loyalty_points // 100
        })


class UserGamingItemListAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        gaming_items = (
            GamingItem.objects.select_related("category")
            .filter(
                is_active=True,
                is_deleted=False,
                maintenance_mode=False,
            )
            .order_by("category__name", "name")
        )

        serializer = UserGamingItemSerializer(
            gaming_items,
            many=True,
            context={"request": request},
        )
        return Response(serializer.data)








from rest_framework.generics import ListAPIView


class UserUpcomingEventsAPIView(ListAPIView):

    serializer_class = UpcomingEventSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return (
            ExclusiveEvent.objects
            .filter(
                is_deleted=False,
                is_active=True,
                event_date__gte=timezone.localdate(),
            )
            .order_by("event_date", "start_time")
        )
    
class UserBookEventAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def post(self, request, pk):

        try:
            event = ExclusiveEvent.objects.get(
                pk=pk,
                is_deleted=False,
                is_active=True
            )
        except ExclusiveEvent.DoesNotExist:
            return Response(
                {"detail": "Event not found."},
                status=404,
            )

        if event.event_date < timezone.localdate():
            return Response(
                {"detail": "Event already finished."},
                status=400,
            )

        if EventBooking.objects.filter(
            event=event,
            user=request.user
        ).exclude(
            status=EventBooking.Status.CANCELLED
        ).exists():

            return Response(
                {"detail": "You have already booked this event."},
                status=400,
            )

        if event.available_slots <= 0:
            return Response(
                {"detail": "Event is full."},
                status=400,
            )

        booking = EventBooking.objects.create(
            event=event,
            user=request.user,
            full_name=request.user.full_name,
            email=request.user.email,
            phone=request.user.phone,
            amount_paid=event.price,
            status=EventBooking.Status.PENDING,
        )

        return Response(
                EventBookingSerializer(
                    booking,
                    context={"request": request},
                ).data,
                status=status.HTTP_201_CREATED,
            )
    


class UserEventBookingsAPIView(ListAPIView):

    serializer_class = EventBookingSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return (
            EventBooking.objects
            .filter(user=self.request.user)
            .select_related("event")
            .order_by("-created_at")
        )

    def get_serializer_context(self):
        return {
            "request": self.request,
        }
    

class UserUpcomingBookingsAPIView(ListAPIView):

    serializer_class = EventBookingSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):

        today = timezone.localdate()

        return (
            EventBooking.objects
            .filter(
                user=self.request.user,
                event__event_date__gte=today,
            )
            .exclude(
                status=EventBooking.Status.CANCELLED
            )
            .select_related("event")
            .order_by("event__event_date")
        )
    


class UserPreviousBookingsAPIView(ListAPIView):

    serializer_class = EventBookingSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):

        today = timezone.localdate()

        return (
            EventBooking.objects
            .filter(
                user=self.request.user,
                event__event_date__lt=today,
            )
            .select_related("event")
            .order_by("-event__event_date")
        )
    


class UserCancelEventBookingAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def post(self, request, pk):

        try:
            booking = EventBooking.objects.select_related("event").get(
                pk=pk,
                user=request.user,
            )
        except EventBooking.DoesNotExist:
            return Response(
                {"detail": "Booking not found."},
                status=404,
            )

        if booking.status == EventBooking.Status.CANCELLED:
            return Response(
                {"detail": "Already cancelled."},
                status=400,
            )

        if booking.event.event_date <= timezone.localdate():
            return Response(
                {"detail": "Cannot cancel after the event date."},
                status=400,
            )

        booking.status = EventBooking.Status.CANCELLED
        booking.save(update_fields=["status"])

        return Response(
            {"detail": "Booking cancelled successfully."}
        )
    











from django.db import transaction
from django.utils import timezone
from rest_framework import status
from rest_framework.generics import ListAPIView, RetrieveAPIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from gaming.models import Booking, ComboPack, LoyaltyTransaction
from gaming.serializers import (
    CreateComboBookingSerializer,
    UserComboBookingSerializer,
    UserComboPackDetailSerializer,
    UserComboPackSerializer,
)


def combo_queryset():
    queryset = ComboPack.objects.filter(is_deleted=False, is_active=True)
    for relation in ("gaming_items", "items"):
        try:
            queryset = queryset.prefetch_related(relation)
            break
        except Exception:
            pass
    return queryset


def booking_queryset(user):
    return (
        Booking.objects.filter(user=user, combo_pack__isnull=False)
        .select_related("combo_pack", "user")
        .prefetch_related("members")
    )


class UserComboPackListAPIView(ListAPIView):
    serializer_class = UserComboPackSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return combo_queryset().order_by("name")

    def get_serializer_context(self):
        return {"request": self.request}


class UserComboPackDetailAPIView(RetrieveAPIView):
    serializer_class = UserComboPackDetailSerializer
    permission_classes = [IsAuthenticated]
    lookup_field = "id"
    lookup_url_kwarg = "id"

    def get_queryset(self):
        return combo_queryset()

    def get_serializer_context(self):
        return {"request": self.request}


class CreateComboBookingAPIView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = CreateComboBookingSerializer(data=request.data, context={"request": request})
        serializer.is_valid(raise_exception=True)
        booking = serializer.save()
        booking = booking_queryset(request.user).get(pk=booking.pk)
        return Response(
            UserComboBookingSerializer(booking, context={"request": request}).data,
            status=status.HTTP_201_CREATED,
        )


class UserComboBookingsAPIView(ListAPIView):
    serializer_class = UserComboBookingSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return booking_queryset(self.request.user).order_by("-created_at")

    def get_serializer_context(self):
        return {"request": self.request}


class UpcomingComboBookingsAPIView(ListAPIView):
    serializer_class = UserComboBookingSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return (
            booking_queryset(self.request.user)
            .filter(booking_date__gte=timezone.localdate())
            .exclude(status=Booking.Status.CANCELLED)
            .order_by("booking_date", "start_time")
        )

    def get_serializer_context(self):
        return {"request": self.request}


class PreviousComboBookingsAPIView(ListAPIView):
    serializer_class = UserComboBookingSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return (
            booking_queryset(self.request.user)
            .filter(booking_date__lt=timezone.localdate())
            .order_by("-booking_date", "-start_time")
        )

    def get_serializer_context(self):
        return {"request": self.request}


class CancelComboBookingAPIView(APIView):
    permission_classes = [IsAuthenticated]

    @transaction.atomic
    def post(self, request, id):
        try:
            booking = booking_queryset(request.user).select_for_update().get(id=id)
        except Booking.DoesNotExist:
            return Response({"detail": "Booking not found."}, status=status.HTTP_404_NOT_FOUND)

        completed_status = getattr(getattr(Booking, "Status", object), "COMPLETED", "COMPLETED")
        if booking.status == completed_status:
            return Response({"detail": "Completed booking cannot be cancelled."}, status=status.HTTP_400_BAD_REQUEST)
        if booking.status == Booking.Status.CANCELLED:
            return Response({"detail": "Booking is already cancelled."}, status=status.HTTP_400_BAD_REQUEST)
        if booking.booking_date <= timezone.localdate():
            return Response({"detail": "Cannot cancel on or after booking date."}, status=status.HTTP_400_BAD_REQUEST)

        booking.status = Booking.Status.CANCELLED
        booking.save(update_fields=["status"])

        points = int(getattr(booking.combo_pack, "loyalty_bonus", 0) or 0)
        if points:
            user = booking.user
            current_points = int(getattr(user, "loyalty_points", 0) or 0)
            user.loyalty_points = max(0, current_points - points)
            user.save(update_fields=["loyalty_points"])
            LoyaltyTransaction.objects.create(
                user=user,
                booking=booking,
                points=-points,
                transaction_type=LoyaltyTransaction.TransactionType.REDEEMED,
                description=f"Combo booking cancelled ({booking.booking_id})",
            )

        return Response(
            {
                "message": "Booking cancelled successfully.",
                "booking": UserComboBookingSerializer(booking, context={"request": request}).data,
            }
        )







from rest_framework.views import APIView
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework import status

from gaming.models import TermsAndConditions
from gaming.serializers import TermsSerializer


class TermsAPIView(APIView):

    permission_classes = [AllowAny]

    def get(self, request):

        terms = (
            TermsAndConditions.objects
            .filter(
                is_current=True,
                is_active=True,
                is_deleted=False,
            )
            .first()
        )

        if not terms:

            return Response(
                {
                    "detail": "Terms & Conditions not found."
                },
                status=status.HTTP_404_NOT_FOUND
            )

        serializer = TermsSerializer(
            terms,
            context={"request": request}
        )

        return Response(serializer.data)










from rest_framework import viewsets
from rest_framework.permissions import IsAuthenticated

from gaming.models import HappyHourAllocation
from gaming.serializers import HappyHourAllocationSerializer
from gaming.permissions import IsAdminUserRole


class AdminHappyHourAllocationViewSet(viewsets.ModelViewSet):

    queryset = HappyHourAllocation.objects.all()
    serializer_class = HappyHourAllocationSerializer

    permission_classes = [
        IsAuthenticated,
        IsAdminUserRole,
    ]

    def get_queryset(self):
        queryset = HappyHourAllocation.objects.all()

        date = self.request.query_params.get("date")
        is_active = self.request.query_params.get("is_active")

        if date:
            queryset = queryset.filter(date=date)

        if is_active is not None:
            queryset = queryset.filter(
                is_active=is_active.lower() == "true"
            )

        return queryset.order_by("date", "start_time")
    






from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status

from gaming.models import HappyHourAllocation
from gaming.serializers import UserHappyHourAllocationSerializer


class UserHappyHourAPIView(APIView):

    permission_classes = [IsAuthenticated]

    def get(self, request):

        today = timezone.localdate()
        current_time = timezone.localtime().time()

        # Get active future/current Happy Hour allocations
        allocations = HappyHourAllocation.objects.filter(
            is_active=True,
            date__gte=today,
        ).order_by(
            "date",
            "start_time",
        )

        upcoming = []

        for allocation in allocations:

            # Future date
            if allocation.date > today:
                upcoming.append(allocation)

            # Today
            elif allocation.date == today:

                # Keep current or upcoming slots
                if allocation.end_time > current_time:
                    upcoming.append(allocation)

        # Only return the nearest Happy Hour
        if not upcoming:
            return Response(
                {
                    "happy_hour": None
                },
                status=status.HTTP_200_OK,
            )

        happy_hour = upcoming[0]

        serializer = UserHappyHourAllocationSerializer(
            happy_hour
        )

        return Response(
            {
                "happy_hour": serializer.data
            },
            status=status.HTTP_200_OK,
        )